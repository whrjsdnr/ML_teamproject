"""Authoritative report renderer. Shared Korean content JSON + verified result files."""
import os
os.environ['MPLCONFIGDIR']='/tmp/final-report-matplotlib'
from pathlib import Path
import json, html, hashlib
import pandas as pd
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager,ticker
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate,Paragraph,Table,TableStyle,Image,Spacer,PageBreak,KeepInFrame
from reportlab.lib.pagesizes import A4
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path.cwd();O=R/'results/final_report';F=O/'figures';D=R/'reports';base='CSE_CIC_IDS2018_Security_AI_Final_Experiment_Report'
font='/usr/share/fonts/truetype/nanum/NanumGothic.ttf';bold='/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf'
font_manager.fontManager.addfont(font);font_manager.fontManager.addfont(bold)
plt.rcParams.update({'font.family':['NanumGothic','DejaVu Sans'],'axes.unicode_minus':False,'mathtext.fontset':'dejavusans','mathtext.default':'regular','font.size':10,'axes.spines.top':False,'axes.spines.right':False})
pages=json.loads((D/(base+'.content.json')).read_text())
N=['lightgbm','fp32','int8'];L=['LightGBM V4-B2','DNN FP32','TorchAO INT8'];C=['#2265A8','#1B977C','#D88832'];navy='#142D4E'
b=pd.concat([pd.read_csv(O/f'benchmark_{n}.csv') for n in N],ignore_index=True)
m=pd.DataFrame([json.loads((O/f'memory_{n}.json').read_text()) for n in N]);b.to_csv(O/'inference_benchmark_verified.csv',index=False);m.to_csv(O/'runtime_memory_verified.csv',index=False)
q=pd.DataFrame([pd.read_csv(R/'results/lightgbm_top40_v4b2/metrics.csv').iloc[0],pd.read_csv(R/'results/dnn_top40_v4/dnn_fp32_full_valid_metrics.csv').iloc[0],pd.read_csv(R/'results/dnn_top40_v4/dnn_torchao_int8_full_valid_metrics.csv').iloc[0]])
files=['models/network/lightgbm_top40_v4b2.txt','results/ondevice_benchmark_v4/dnn_fp32_state_dict.pt','results/ondevice_benchmark_v4/dnn_torchao_int8_state_dict.pt'];sizes=[(R/x).stat().st_size/1024**2 for x in files]
comp=pd.DataFrame({'model':L,'macro_f1':q.macro_f1.to_numpy(),'accuracy':q.accuracy.to_numpy(),'serialized_size_mib':sizes,'model_file':files})
for bs in [1,32,256,4096]:comp[f'batch{bs}_median_latency_ms']=[b[(b.model==n)&(b.batch_size==bs)].iloc[0].median_latency_ms for n in N]
comp['batch4096_throughput_flows_s']=[b[(b.model==n)&(b.batch_size==4096)].iloc[0].throughput_flows_s for n in N]
for col in m.columns:
 if col!='model':comp[col]=m[col].to_numpy()
comp['benchmark_scope']='report remeasurement; WSL2 x86 CPU 1-thread; 2 warmups; LGBM 3 / DNN 30 repeats; not embedded hardware';comp.to_csv(O/'ondevice_model_comparison.csv',index=False)
def fig(name,title,ylabel,draw,log=False):
 f,ax=plt.subplots(figsize=(8,3.75));draw(ax);ax.set_title(title,loc='left',color=navy,fontweight='bold',pad=13);ax.set_ylabel(ylabel)
 if log:ax.set_yscale('log');ax.yaxis.set_major_formatter(ticker.FuncFormatter(lambda x,pos:f'{x:g}'));ax.set_ylabel(ylabel+' (log scale)')
 ax.grid(axis='y',alpha=.18);f.tight_layout();f.savefig(F/name,dpi=220,bbox_inches='tight');plt.close(f)
def bar(name,title,vals,ylabel,log=False):
 def draw(ax):
  ax.bar(L,vals,color=C)
  for i,v in enumerate(vals):ax.annotate(f'{v:.4f}',(i,v),xytext=(0,4),textcoords='offset points',ha='center',fontsize=9)
  ax.margins(y=.22)
 fig(name,title,ylabel,draw,log)
bar('02_model_macro_f1.png','동일 Full Validation: 탐지 품질 비교',comp.macro_f1,'Macro F1 (0–1)')
bar('03_model_size.png','추론 가중치 / 모델 파일 크기',sizes,'Serialized file size (MiB)',True)
bar('06_model_load_rss.png','모델 로드 전후 process RSS 순변화',m.model_load_delta_mb,'RSS delta (MiB)')
bar('07_peak_delta_rss.png','프로세스 baseline 대비 peak RSS 증가',m.peak_delta_from_baseline_mb,'Peak RSS delta (MiB)')
for field,name,title,y in [('median_latency_ms','04_latency_by_batch.png','Batch size별 순수 추론 median latency','Latency (ms / batch)'),('throughput_flows_s','05_throughput_by_batch.png','Batch size별 추론 처리량','Throughput (flows/s)')]:
 def draw(ax,field=field):
  for n,l,c in zip(N,L,C):
   bb=b[b.model==n];ax.plot(bb.batch_size.astype(str),bb[field],marker='o',color=c,label=l)
  ax.set_xlabel('Batch size (flows)');ax.legend(fontsize=9)
 fig(name,title,y,draw,True)
mw=pd.read_csv(R/'results/statistical_analysis_v4/mann_whitney_top40.csv')
def effect(ax):
 d=mw.head(10).iloc[::-1];ax.barh(d.feature,d.effect_size,color=[C[0] if v<0 else C[1] for v in d.effect_size]);ax.axvline(0,color=navy,lw=.7);ax.set_xlabel('Rank-biserial effect size: 2U/(nB*nBot) - 1');ax.set_ylabel('Feature')
fig('10_statistical_effect_size.png','Benign vs Bot: 주요 signed effect size','',effect)
# Refresh report tables from current authoritative verification files.
for pg in pages:
 for block in pg['blocks']:
  if block['kind']=='table':
   cap=block['caption']
   if cap=='재측정 순수 inference latency (ms)':block['data']=[['모델','batch1','batch32','batch256','batch4096']]+[[L[i]]+[f'{comp.iloc[i][f"batch{bs}_median_latency_ms"]:.4f}' for bs in [1,32,256,4096]] for i in range(3)]
   if cap=='Batch4096 throughput':block['data']=[['모델','실제 재측정 flows/s','제공 참고 flows/s']]+[[L[i],f'{comp.iloc[i].batch4096_throughput_flows_s:,.0f}',['574','1,253,030','64,468'][i]] for i in range(3)]
   if cap=='Fresh Python process의 재측정 RSS (MiB)':block['data']=[['Stage','LightGBM','FP32','INT8']]+[[col]+[f'{m.iloc[i][col]:.4f}' for i in range(3)] for col in m.columns if col!='model']
   if cap.startswith('Test V2:') or cap in ['FP32 Full Validation','TorchAO - FP32: 같은 Full Validation']:
    if block['data'][0][0] not in ['Metric','지표']:block['data'].insert(0,['Metric','Value'])
  if block['kind'] in ['p','h','source']:block['text']=block['text'].replace('−','-').replace('build_report.py','render_final_report.py')
  if block['kind']=='figure':block['caption']=block['caption'].replace('−','-')
(D/(base+'.content.json')).write_text(json.dumps(pages,ensure_ascii=False,indent=2))
# Verify the scope of statistical sample from labels only.
import pyarrow.parquet as pq
stat=pq.read_table(R/'results/statistical_analysis_v4/benign_attack_sample.parquet',columns=['label_id','binary_label']).to_pandas()
stat.value_counts().rename('rows').reset_index().to_csv(O/'statistical_sample_label_audit.csv',index=False)
# Verified metric summary with provenance.
rows=[]
def add(df,src,scope):
 for ix,r in df.iterrows():
  for k,v in r.items():
   if isinstance(v,(int,float,np.number)) and pd.notna(v):rows.append(dict(model=r.get('model',str(ix)),scope=scope,metric=k,value=float(v),source=src,status='verified'))
for src,scope in [('results/model_comparison_v3/model_diversity_comparison.csv','V3 sampled validation'),('results/xgboost_v3/validation_metrics.csv','V3 Full78 sampled validation'),('results/feature_selection/feature_reduction_comparison_v3.csv','V3 feature reduction'),('results/data_quality/dataset_profile.csv','original EDA dataset profile')]:add(pd.read_csv(R/src),src,scope)
for n in ['v4a','v4b1','v4b2']:
 src=f'results/lightgbm_top40_{n}/metrics.csv';add(pd.read_csv(R/src),src,'V4 Full Validation')
for n in ['fp32','torchao_int8']:
 src=f'results/dnn_top40_v4/dnn_{n}_full_valid_metrics.csv';df=pd.read_csv(R/src);df['model']='DNN '+n;add(df,src,'V4 Full Validation')
cm=pd.read_csv(R/'results/final_evaluation/test_confusion_matrix.csv').iloc[:,1:].to_numpy();tp=cm.diagonal();sup=cm.sum(1);p=np.divide(tp,cm.sum(0),out=np.zeros(15),where=cm.sum(0)>0);rr=tp/sup;ff=np.divide(2*p*rr,p+rr,out=np.zeros(15),where=p+rr>0);test={'model':'LightGBM V3','accuracy':tp.sum()/cm.sum(),'macro_precision':p.mean(),'macro_recall':rr.mean(),'macro_f1':ff.mean(),'weighted_f1':np.dot(ff,sup)/cm.sum()};add(pd.DataFrame([test]),'results/final_evaluation/test_confusion_matrix.csv','consumed Test V2; audit only')
add(comp,'results/final_report/ondevice_model_comparison.csv','WSL CPU report remeasurement')
audit=json.loads((O/'data_verification.json').read_text())
for key in ['conflicting_patterns','affected_rows','max_labels','0_1_overlap','0_2_overlap','1_2_overlap']:
 rows.append(dict(model='78-feature hash audit',scope='entire Split V2',metric=key,value=audit[key],source='results/final_report/data_verification.json',status='verified'))
for split,count in audit['rows'].items():rows.append(dict(model=split,scope='Split V2 Parquet footer',metric='rows',value=count,source='data/processed/network_ml_split_v2/'+split,status='verified'))
for path,scope in [('results/final_report/label_conflict_pairs_verified.csv','entire Split V2 conflict pairs'),('results/statistical_analysis_v4/mann_whitney_top40.csv','Benign vs Bot; LIMIT sample'),('results/statistical_analysis_v4/kruskal_wallis_top40.csv','15 classes; LIMIT sample')]:
 df=pd.read_csv(R/path)
 if 'feature' in df:df['model']=df.feature
 elif 'class_a' in df:df['model']=df.class_a.astype(str)+' vs '+df.class_b.astype(str)
 add(df,path,scope)
for key,i,j in [('Benign_to_Infilteration',0,12),('SlowHTTP_to_FTP',9,11),('Infilteration_to_Benign',12,0),('FTP_to_SlowHTTP',11,9)]:rows.append(dict(model='LightGBM V3',scope='consumed Test V2 audit',metric=key,value=int(cm[i,j]),source='results/final_evaluation/test_confusion_matrix.csv',status='verified'))
for key,val in [('MW_FDR_significant',int(mw.significant_fdr_005.sum())),('KW_FDR_significant',int(pd.read_csv(R/'results/statistical_analysis_v4/kruskal_wallis_top40.csv').significant_fdr_005.sum()))]:rows.append(dict(model='Top40 statistical analysis',scope='Train samples; MW Bot-only',metric=key,value=val,source='results/statistical_analysis_v4/',status='verified'))
pd.DataFrame(rows).to_csv(O/'final_metrics_summary.csv',index=False)
# Korean typography and page-fit layout.
pdfmetrics.registerFont(TTFont('Korean',font));pdfmetrics.registerFont(TTFont('KoreanBold',bold))
body=ParagraphStyle('Body',fontName='Korean',fontSize=9,leading=14.5,spaceAfter=9,wordWrap='CJK',textColor=colors.HexColor('#24354A'))
heading=ParagraphStyle('Heading',parent=body,fontName='KoreanBold',fontSize=11,leading=17,spaceAfter=7,textColor=colors.HexColor(navy))
title=ParagraphStyle('Title',parent=heading,fontSize=21,leading=30,spaceAfter=18)
small=ParagraphStyle('Small',parent=body,fontSize=7.2,leading=11,spaceAfter=5,textColor=colors.HexColor('#59697B'))
cell=ParagraphStyle('Cell',parent=body,fontSize=8,leading=12,spaceAfter=0)
head=ParagraphStyle('Head',parent=cell,fontName='KoreanBold',textColor=colors.white)
def para(s,style=body):return Paragraph(html.escape(str(s)).replace('\n','<br/>'),style)
story=[];frames=[];doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929);sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.72)
for sn in ['Normal','Title','Heading 1','Heading 2']:
 st=doc.styles[sn];st.font.name='나눔고딕';st._element.rPr.rFonts.set(qn('w:eastAsia'),'나눔고딕')
 st.font.size=Pt(9 if sn=='Normal' else 20 if sn in ['Title','Heading 1'] else 11)
 st.font.color.rgb=RGBColor.from_string('142D4E' if sn!='Normal' else '24354A')
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
footer=sec.footer.paragraphs[0];footer.alignment=2;footer.add_run('SECURITY AI | ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
fig_no=table_no=0
for idx,pg in enumerate(pages):
 if idx:story.append(PageBreak());doc.add_page_break()
 content=[para(pg['kicker'],small),para(pg['title'],title)];doc.add_paragraph(pg['kicker']);doc.add_heading(pg['title'],0 if idx==0 else 1)
 for block in pg['blocks']:
  k=block['kind']
  if k in ['p','h','source']:
   content.append(para(block['text'],{'p':body,'h':heading,'source':small}[k]));dp=doc.add_heading(block['text'],2) if k=='h' else doc.add_paragraph(block['text'])
   if k=='source':
    for run in dp.runs:run.font.size=Pt(7)
  elif k=='table':
   table_no+=1;cap=f'표 {table_no}. '+block['caption'];content.append(para(cap,small));data=block['data'];widths=block.get('widths') or [485/len(data[0])]*len(data[0]);cells=[[para(x,head if j==0 else cell) for x in row] for j,row in enumerate(data)]
   t=Table(cells,colWidths=widths,repeatRows=1,hAlign='LEFT');t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor(navy)),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.HexColor('#F0F4F9'),colors.white]),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4),('LINEBELOW',(0,-1),(-1,-1),.5,colors.HexColor('#D1DCE8'))]));content.extend([t,Spacer(1,10)])
   doc.add_paragraph(cap);dt=doc.add_table(rows=1,cols=len(data[0]));dt.style='Light Shading Accent 1';tr=dt.rows[0]._tr.get_or_add_trPr();hdr=OxmlElement('w:tblHeader');tr.append(hdr)
   for i,x in enumerate(data[0]):dt.rows[0].cells[i].text=str(x)
   for row in data[1:]:
    cc=dt.add_row().cells
    for i,x in enumerate(row):cc[i].text=str(x)
   for row in dt.rows:
    for c in row.cells:
     for pp in c.paragraphs:
      for rr in pp.runs:rr.font.size=Pt(8)
  elif k=='figure':
   fig_no+=1;im=Image(str(R/block['path']));ratio=im.imageWidth/im.imageHeight;hh=min(block['height'],485/ratio);im.drawHeight=hh;im.drawWidth=hh*ratio;im.hAlign='LEFT';cap=f'그림 {fig_no}. '+block['caption'];content.extend([im,Spacer(1,4),para(cap,small),Spacer(1,6)]);doc.add_picture(str(R/block['path']),width=Inches(min(6.7,hh*ratio/72)));doc.add_paragraph(cap)
 frame=KeepInFrame(485,730,content,mode='shrink',hAlign='LEFT',vAlign='TOP');frames.append(frame);story.append(frame)
def footerpdf(canvas,document):
 canvas.setStrokeColor(colors.HexColor('#D1DCE8'));canvas.line(50,42,A4[0]-50,42);canvas.setFont('Korean',7);canvas.setFillColor(colors.HexColor('#59697B'));canvas.drawString(50,29,'CSE-CIC-IDS2018 | Security AI | 저장소 검증 기반');canvas.drawRightString(A4[0]-50,29,str(document.page))
SimpleDocTemplate(str(D/(base+'.pdf')),pagesize=A4,rightMargin=50,leftMargin=50,topMargin=44,bottomMargin=53,title='CSE-CIC-IDS2018 Security AI 최종 실험 보고서').build(story,onFirstPage=footerpdf,onLaterPages=footerpdf)
doc.save(str(D/(base+'.docx')));(O/'layout_scales.json').write_text(json.dumps([getattr(x,'_scale',1) for x in frames]))
# Manifest with explicit file-level evidence.
sources=set(files+['README.md','notebooks/preprocessing.ipynb','notebooks/model_sample_v4.ipynb','notebooks/statistical_analysis_v4.ipynb','notebooks/final_evaluation.ipynb','notebooks/dnn_top40_v4.ipynb','notebooks/ondevice_benchmark_v4.ipynb','scripts/benchmark_model_memory.py','models/network/dnn_top40_v4_scaler.joblib','models/network/dnn_top40_v4_best.pt','models/network/dnn_top40_v4_dynamic_int8.pt','results/statistical_analysis_v4/benign_attack_sample.parquet'])
for pat in ['results/data_quality/*.csv','results/network_eda/*.csv','results/statistical_analysis_v4/*.csv','results/model_sample_v4/*.csv','results/v4_model_comparison/*.csv','results/dnn_top40_v4/*.csv','results/feature_selection/*.csv','results/model_comparison_v3/*.csv','results/final_evaluation/*.csv','results/final_report/*verified*.csv','results/final_report/memory_*.json','results/final_report/benchmark*.csv']:
 sources.update(str(p.relative_to(R)) for p in R.glob(pat))
for pg in pages:
 for block in pg['blocks']:
  if block['kind']=='figure':sources.add(block['path'])
manifest=[]
for p in sorted(sources):
 pp=R/p;h=hashlib.sha256()
 with pp.open('rb') as fin:
  for chunk in iter(lambda:fin.read(1024*1024),b''):h.update(chunk)
 manifest.append({'path':p,'bytes':pp.stat().st_size,'sha256':h.hexdigest()})
pd.DataFrame(manifest).to_csv(O/'source_manifest.csv',index=False)
print('PDF/DOCX generated:',len(pages),'pages;',fig_no,'figures;',table_no,'tables; max shrink',max(getattr(x,'_scale',1) for x in frames))
